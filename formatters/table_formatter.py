from .base_formatter import BaseFormatter
from utils.docx_helper import (
    set_paragraph_format, format_paragraph_text, match_pattern,
    set_cell_border, set_cell_vertical_alignment, set_run_font,
    split_text_by_language,
)
from docx.enum.text import WD_ALIGN_PARAGRAPH


class TableFormatter(BaseFormatter):
    def __init__(self, doc, config):
        super().__init__(doc, config)
        self.table_config = config.get('table', {})
        self.caption_config = self.table_config.get('caption', {})
        self.content_config = self.table_config.get('content', {})
        self.style_config = self.table_config.get('style', {})
        self.table_pattern = self.recognition_config.get('table_caption_pattern', '')

    def format(self):
        self._format_table_captions()
        for table in self.doc.tables:
            self._format_table(table)

    def _format_table_captions(self):
        for paragraph in self.doc.paragraphs:
            if match_pattern(paragraph.text, self.table_pattern):
                self._format_caption(paragraph)

    def _format_caption(self, paragraph):
        cfg = self.caption_config

        set_paragraph_format(
            paragraph,
            alignment=cfg.get('alignment', 'center'),
            space_before=cfg.get('space_before', 6),
            space_after=cfg.get('space_after', 6),
        )

        format_paragraph_text(
            paragraph,
            cn_font=cfg.get('font_name_cn', '宋体'),
            en_font=cfg.get('font_name_en', 'Times New Roman'),
            font_size=cfg.get('font_size', 12),
            bold=cfg.get('bold', True),
        )

    def _format_table(self, table):
        if self.style_config.get('border_style') == 'three_line':
            self._apply_three_line_style(table)

        for row in table.rows:
            for cell in row.cells:
                self._format_cell(cell)

    def _apply_three_line_style(self, table):
        top_w = int(self.style_config.get('top_border_width', 1.5) * 8)
        mid_w = int(self.style_config.get('middle_border_width', 0.75) * 8)
        bot_w = int(self.style_config.get('bottom_border_width', 1.5) * 8)

        num_rows = len(table.rows)

        for i, row in enumerate(table.rows):
            for cell in row.cells:
                borders = {}

                if i == 0:
                    borders['top'] = {'sz': top_w, 'val': 'single', 'color': '000000'}
                    borders['bottom'] = {'sz': mid_w, 'val': 'single', 'color': '000000'}
                elif i == num_rows - 1:
                    borders['top'] = {'sz': 0, 'val': 'none', 'color': '000000'}
                    borders['bottom'] = {'sz': bot_w, 'val': 'single', 'color': '000000'}
                else:
                    borders['top'] = {'sz': 0, 'val': 'none', 'color': '000000'}
                    borders['bottom'] = {'sz': 0, 'val': 'none', 'color': '000000'}

                borders['left'] = {'sz': 0, 'val': 'none', 'color': '000000'}
                borders['right'] = {'sz': 0, 'val': 'none', 'color': '000000'}

                set_cell_border(cell, **borders)

    def _format_cell(self, cell):
        cfg = self.content_config

        set_cell_vertical_alignment(cell, cfg.get('vertical_alignment', 'center'))

        align_map = {
            'left': WD_ALIGN_PARAGRAPH.LEFT,
            'center': WD_ALIGN_PARAGRAPH.CENTER,
            'right': WD_ALIGN_PARAGRAPH.RIGHT,
        }
        target_align = align_map.get(cfg.get('alignment', 'center'), WD_ALIGN_PARAGRAPH.CENTER)

        for paragraph in cell.paragraphs:
            paragraph.alignment = target_align

            text = paragraph.text
            if not text.strip():
                continue

            paragraph.clear()
            segments = split_text_by_language(text)

            for seg_text, is_en in segments:
                run = paragraph.add_run(seg_text)
                font_name = self.get_english_font() if is_en else cfg.get('font_name_cn', '宋体')
                set_run_font(run, font_name, cfg.get('font_size', 12), cfg.get('bold', False))
