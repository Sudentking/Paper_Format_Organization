import re
import base64
import os
from html import escape

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

PREVIEW_CSS = """
<style>
    @page { size: A4; margin: 2.54cm; }
    body {
        font-family: "宋体", "SimSun", "Times New Roman", serif;
        font-size: 12pt;
        line-height: 1.5;
        color: #333;
        max-width: 210mm;
        margin: 0 auto;
        padding: 20px 40px;
        background: white;
    }
    h1, .heading-1 {
        font-family: "黑体", "SimHei", sans-serif;
        font-size: 16pt;
        font-weight: bold;
        text-align: center;
        margin: 24px 0 12px;
        line-height: 1.5;
    }
    h2, .heading-2 {
        font-family: "黑体", "SimHei", sans-serif;
        font-size: 14pt;
        font-weight: bold;
        margin: 18px 0 8px;
        line-height: 1.5;
    }
    h3, .heading-3 {
        font-family: "黑体", "SimHei", sans-serif;
        font-size: 12pt;
        font-weight: bold;
        margin: 12px 0 6px;
        line-height: 1.5;
    }
    p {
        margin: 6px 0;
        line-height: 1.5;
    }
    table {
        border-collapse: collapse;
        width: 100%;
        margin: 12px 0;
        font-size: 10.5pt;
    }
    th, td {
        border: 1px solid #333;
        padding: 6px 10px;
        text-align: center;
        vertical-align: middle;
    }
    th {
        font-weight: bold;
        background: #f5f5f5;
    }
    img {
        max-width: 100%;
        height: auto;
        display: block;
        margin: 12px auto;
    }
    a { color: #1a73e8; text-decoration: none; }
    a:hover { text-decoration: underline; }
    .toc-entry { margin: 2px 0; }
    .toc-entry a { color: #333; text-decoration: none; }
</style>
"""

PREVIEW_WARNING = """
<div style="background:#fff3cd;border:1px solid #ffc107;border-radius:6px;padding:10px 16px;margin-bottom:16px;font-size:13px;color:#856404;">
    ⚠️ 预览为 HTML 渲染效果，可能与最终 Word 文档存在细微差异（字体、分页、页眉页脚等）。
    如需精确排版效果，请下载 .docx 文件后在 Word 中查看。
</div>
"""


def _escape(text):
    if not text:
        return ''
    return escape(text)


def _run_to_html(run):
    text = run.text
    if not text:
        return ''
    text = _escape(text)
    parts = []
    style_parts = []

    if run.bold:
        parts.append('<strong>')
    if run.italic:
        parts.append('<em>')
    if run.underline:
        parts.append('<u>')
    if run.font.strike:
        parts.append('<s>')

    if run.font.size:
        size_pt = run.font.size.pt
        style_parts.append(f'font-size:{size_pt}pt')
    if run.font.name:
        style_parts.append(f'font-family:"{run.font.name}",serif')
    if run.font.color and run.font.color.rgb:
        style_parts.append(f'color:#{run.font.color.rgb}')

    if style_parts:
        text = f'<span style="{";".join(style_parts)}">{text}</span>'

    if run.font.strike:
        text += '</s>'
    if run.underline:
        text += '</u>'
    if run.italic:
        text += '</em>'
    if run.bold:
        text += '</strong>'

    return text


def _detect_heading_level(paragraph):
    pPr = paragraph._element.find(qn('w:pPr'))
    if pPr is not None:
        outline_lvl = pPr.find(qn('w:outlineLvl'))
        if outline_lvl is not None:
            val = outline_lvl.get(qn('w:val'))
            if val is not None:
                try:
                    return int(val) + 1
                except (ValueError, TypeError):
                    pass

    style_name = paragraph.style.name if paragraph.style else ''
    if style_name.startswith('Heading'):
        level_match = re.search(r'(\d+)', style_name)
        if level_match:
            return int(level_match.group(1))
    elif 'heading' in style_name.lower() or '标题' in style_name:
        if '1' in style_name or '一' in style_name:
            return 1
        elif '2' in style_name or '二' in style_name:
            return 2
        elif '3' in style_name or '三' in style_name:
            return 3

    return None


def _paragraph_to_html(paragraph):
    text = paragraph.text
    if not text and not paragraph.runs:
        return ''

    style_parts = []
    class_parts = []
    tag = 'p'

    heading_level = _detect_heading_level(paragraph)
    if heading_level is not None:
        tag = f'h{min(heading_level, 6)}'

    align = paragraph.alignment
    if align is not None:
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        align_map = {
            WD_ALIGN_PARAGRAPH.CENTER: 'center',
            WD_ALIGN_PARAGRAPH.RIGHT: 'right',
            WD_ALIGN_PARAGRAPH.JUSTIFY: 'justify',
            WD_ALIGN_PARAGRAPH.LEFT: 'left',
        }
        css_align = align_map.get(align)
        if css_align:
            style_parts.append(f'text-align:{css_align}')

    pf = paragraph.paragraph_format
    if pf.line_spacing:
        style_parts.append(f'line-height:{pf.line_spacing}')
    if pf.space_before:
        style_parts.append(f'margin-top:{pf.space_before.pt}pt')
    if pf.space_after:
        style_parts.append(f'margin-bottom:{pf.space_after.pt}pt')
    if pf.first_line_indent:
        style_parts.append(f'text-indent:{pf.first_line_indent.pt}pt')
    if pf.left_indent:
        style_parts.append(f'margin-left:{pf.left_indent.pt}pt')

    attrs = ''
    if style_parts:
        attrs += f' style="{";".join(style_parts)}"'
    if class_parts:
        attrs += f' class="{" ".join(class_parts)}"'

    runs_html = ''.join(_run_to_html(r) for r in paragraph.runs)
    if not runs_html.strip():
        runs_html = _escape(text)

    return f'<{tag}{attrs}>{runs_html}</{tag}>'


def _table_to_html(table):
    rows_html = []
    for i, row in enumerate(table.rows):
        cells_html = []
        for cell in row.cells:
            cell_tag = 'th' if i == 0 else 'td'
            cell_text = _escape(cell.text.strip())
            cells_html.append(f'<{cell_tag}>{cell_text}</{cell_tag}>')
        rows_html.append(f'<tr>{"".join(cells_html)}</tr>')

    return f'<table>{"".join(rows_html)}</table>'


def _convert_with_docx(docx_path: str) -> str:
    doc = Document(docx_path)
    body = doc.element.body

    html_parts = []
    for child in body:
        tag = etree.QName(child).localname

        if tag == 'p':
            try:
                paragraph = doc.paragraphs[0]
                for p in doc.paragraphs:
                    if p._element is child:
                        paragraph = p
                        break
                else:
                    from docx.text.paragraph import Paragraph
                    paragraph = Paragraph(child, doc)
                html_parts.append(_paragraph_to_html(paragraph))
            except Exception:
                pass

        elif tag == 'tbl':
            try:
                for t in doc.tables:
                    if t._tbl is child:
                        html_parts.append(_table_to_html(t))
                        break
            except Exception:
                pass

        elif tag == 'sdt':
            try:
                sdt_content = child.find(qn('w:sdtContent'))
                if sdt_content is not None:
                    for p_elem in sdt_content.findall(qn('w:p')):
                        from docx.text.paragraph import Paragraph
                        para = Paragraph(p_elem, doc)
                        html_parts.append(_paragraph_to_html(para))
                    for tbl_elem in sdt_content.findall(qn('w:tbl')):
                        from docx.table import Table
                        table = Table(tbl_elem, doc)
                        html_parts.append(_table_to_html(table))
            except Exception:
                pass

    return '\n'.join(html_parts)


def convert_docx_to_html(docx_path: str) -> str:
    body_html = _convert_with_docx(docx_path)

    full_html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>文档预览</title>
{PREVIEW_CSS}
</head>
<body>
{PREVIEW_WARNING}
{body_html}
</body>
</html>"""

    return full_html
