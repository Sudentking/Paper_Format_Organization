import sys
import os
import unittest
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

from docx import Document
from services.doc_preview import convert_docx_to_html


class TestDocPreview(unittest.TestCase):
    def setUp(self):
        self.doc = Document()
        self.doc.add_heading("测试标题", level=1)
        self.doc.add_paragraph("这是一段测试正文。")
        self.doc.add_heading("二级标题", level=2)
        self.doc.add_paragraph("这是二级标题下的内容。")
        table = self.doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "A1"
        table.cell(0, 1).text = "B1"
        table.cell(1, 0).text = "A2"
        table.cell(1, 1).text = "B2"

        self.tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
        self.doc.save(self.tmp.name)
        self.tmp.close()

    def tearDown(self):
        os.unlink(self.tmp.name)

    def test_convert_returns_html(self):
        html = convert_docx_to_html(self.tmp.name)
        self.assertIn("<!DOCTYPE html>", html)
        self.assertIn("<html", html)
        self.assertIn("</html>", html)
        self.assertIn("<body>", html)
        self.assertIn("</body>", html)

    def test_convert_contains_content(self):
        html = convert_docx_to_html(self.tmp.name)
        self.assertIn("测试标题", html)
        self.assertIn("测试正文", html)
        self.assertIn("二级标题", html)

    def test_convert_contains_css(self):
        html = convert_docx_to_html(self.tmp.name)
        self.assertIn("<style>", html)
        self.assertIn("font-family", html)
        self.assertIn("宋体", html)
        self.assertIn("黑体", html)

    def test_convert_contains_table(self):
        html = convert_docx_to_html(self.tmp.name)
        self.assertIn("<table", html)
        self.assertIn("A1", html)
        self.assertIn("B2", html)

    def test_convert_contains_warning(self):
        html = convert_docx_to_html(self.tmp.name)
        self.assertIn("预览为 HTML 渲染效果", html)

    def test_convert_empty_doc(self):
        empty_doc = Document()
        tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
        empty_doc.save(tmp.name)
        tmp.close()
        try:
            html = convert_docx_to_html(tmp.name)
            self.assertIn("<!DOCTYPE html>", html)
            self.assertIn("<body>", html)
        finally:
            os.unlink(tmp.name)

    def test_convert_nonexistent_file(self):
        with self.assertRaises(Exception):
            convert_docx_to_html("nonexistent_file.docx")


if __name__ == '__main__':
    unittest.main()
