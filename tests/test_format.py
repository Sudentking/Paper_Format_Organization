import sys
import os
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

from docx import Document
from utils.docx_helper import (
    is_english_or_number,
    split_text_by_language,
    match_pattern,
    set_run_font,
    set_paragraph_format,
    format_paragraph_text,
    set_cell_border,
    set_cell_vertical_alignment,
    FONT_SIZE_MAP,
)
from formatters import (
    HeadingFormatter,
    ParagraphFormatter,
    TableFormatter,
    CaptionFormatter,
    ReferenceFormatter,
    TOCFormatter,
)


class TestDocxHelper(unittest.TestCase):
    def test_is_english_or_number(self):
        self.assertTrue(is_english_or_number('a'))
        self.assertTrue(is_english_or_number('Z'))
        self.assertTrue(is_english_or_number('0'))
        self.assertTrue(is_english_or_number('9'))
        self.assertFalse(is_english_or_number('中'))
        self.assertFalse(is_english_or_number('。'))
        self.assertFalse(is_english_or_number(' '))

    def test_split_text_by_language(self):
        result = split_text_by_language("第1章 引言abc123")
        self.assertEqual(len(result), 4)
        self.assertEqual(result[0], ('第', False))
        self.assertEqual(result[1], ('1', True))
        self.assertEqual(result[2], ('章 引言', False))
        self.assertEqual(result[3], ('abc123', True))

    def test_split_empty(self):
        self.assertEqual(split_text_by_language(""), [])
        self.assertEqual(split_text_by_language(None), [])

    def test_match_pattern(self):
        self.assertTrue(match_pattern("1 标题", r'^[0-9]+\s+'))
        self.assertTrue(match_pattern("1.1 标题", r'^[0-9]+\.[0-9]+\s+'))
        self.assertFalse(match_pattern("标题1", r'^[0-9]+\s+'))
        self.assertFalse(match_pattern("", r'^[0-9]+'))
        self.assertFalse(match_pattern("test", ""))

    def test_font_size_map(self):
        self.assertEqual(FONT_SIZE_MAP['四号'], 14)
        self.assertEqual(FONT_SIZE_MAP['小四'], 12)
        self.assertEqual(FONT_SIZE_MAP['五号'], 10.5)
        self.assertEqual(FONT_SIZE_MAP['三号'], 16)


class TestHeadingFormatter(unittest.TestCase):
    def setUp(self):
        self.doc = Document()
        self.config = {
            'global': {'english_number_font': 'Times New Roman'},
            'headings': {
                'level1': {
                    'font_name_cn': '黑体', 'font_name_en': 'Times New Roman',
                    'font_size': 14, 'bold': True, 'alignment': 'left',
                    'space_before': 0, 'space_after': 0,
                },
                'level2': {
                    'font_name_cn': '黑体', 'font_name_en': 'Times New Roman',
                    'font_size': 12, 'bold': True, 'alignment': 'left',
                    'space_before': 0, 'space_after': 0,
                },
                'level3': {
                    'font_name_cn': '宋体', 'font_name_en': 'Times New Roman',
                    'font_size': 12, 'bold': True, 'alignment': 'left',
                    'space_before': 0, 'space_after': 0,
                },
            },
            'recognition': {
                'heading_patterns': {
                    'level1': r'^[0-9]+\s+',
                    'level2': r'^[0-9]+\.[0-9]+\s+',
                    'level3': r'^[0-9]+\.[0-9]+\.[0-9]+\s+',
                },
            },
        }

    def test_detect_heading_level(self):
        from utils.docx_helper import detect_heading_level
        patterns = self.config['recognition']['heading_patterns']
        p1 = self.doc.add_paragraph("1 引言")
        p2 = self.doc.add_paragraph("1.1 背景")
        p3 = self.doc.add_paragraph("1.1.1 详细")
        p4 = self.doc.add_paragraph("这是正文")
        self.assertEqual(detect_heading_level(p1, patterns), 'level1')
        self.assertEqual(detect_heading_level(p2, patterns), 'level2')
        self.assertEqual(detect_heading_level(p3, patterns), 'level3')
        self.assertIsNone(detect_heading_level(p4, patterns))

    def test_format_heading(self):
        self.doc.add_paragraph("1 引言")
        self.doc.add_paragraph("这是正文")
        self.doc.add_paragraph("1.1 背景")

        fmt = HeadingFormatter(self.doc, self.config)
        fmt.format()

        self.assertEqual(self.doc.paragraphs[0].runs[0].font.size.pt, 14)
        self.assertTrue(self.doc.paragraphs[0].runs[0].font.bold)

    def test_heading_with_english(self):
        self.doc.add_paragraph("1 Introduction")
        fmt = HeadingFormatter(self.doc, self.config)
        fmt.format()

        runs = self.doc.paragraphs[0].runs
        self.assertTrue(len(runs) > 0)

    def test_heading_by_word_style(self):
        from utils.docx_helper import detect_heading_level
        patterns = self.config['recognition']['heading_patterns']
        p = self.doc.add_paragraph("引言")
        p.style = self.doc.styles['Heading 1']
        self.assertEqual(detect_heading_level(p, patterns), 'level1')

        p2 = self.doc.add_paragraph("背景")
        p2.style = self.doc.styles['Heading 2']
        self.assertEqual(detect_heading_level(p2, patterns), 'level2')


class TestParagraphFormatter(unittest.TestCase):
    def setUp(self):
        self.doc = Document()
        self.config = {
            'global': {'english_number_font': 'Times New Roman'},
            'paragraph': {
                'font_name_cn': '宋体', 'font_name_en': 'Times New Roman',
                'font_size': 12, 'bold': False, 'line_spacing': 1.5,
                'alignment': 'justify', 'first_line_indent': 0,
                'space_before': 0, 'space_after': 0,
            },
            'recognition': {
                'heading_patterns': {
                    'level1': r'^[0-9]+\s+',
                    'level2': r'^[0-9]+\.[0-9]+\s+',
                    'level3': r'^[0-9]+\.[0-9]+\.[0-9]+\s+',
                },
                'table_caption_pattern': r'^表[0-9]+\.?[0-9]*\s+',
                'figure_caption_pattern': r'^图[0-9]+\.?[0-9]*\s+',
                'reference_title_pattern': r'^参考文献$',
            },
        }

    def test_should_skip(self):
        fmt = ParagraphFormatter(self.doc, self.config)
        self.assertTrue(fmt._should_skip(self.doc.add_paragraph("")))
        self.assertTrue(fmt._should_skip(self.doc.add_paragraph("1 标题")))
        self.assertTrue(fmt._should_skip(self.doc.add_paragraph("表1.1 名称")))
        self.assertTrue(fmt._should_skip(self.doc.add_paragraph("图1 名称")))
        self.assertTrue(fmt._should_skip(self.doc.add_paragraph("参考文献")))
        self.assertFalse(fmt._should_skip(self.doc.add_paragraph("这是正文")))

    def test_format_paragraph(self):
        self.doc.add_paragraph("这是正文段落。")
        fmt = ParagraphFormatter(self.doc, self.config)
        fmt.format()

        para = self.doc.paragraphs[0]
        self.assertEqual(para.runs[0].font.size.pt, 12)


class TestTableFormatter(unittest.TestCase):
    def setUp(self):
        self.doc = Document()
        self.config = {
            'global': {'english_number_font': 'Times New Roman'},
            'table': {
                'caption': {
                    'font_name_cn': '宋体', 'font_name_en': 'Times New Roman',
                    'font_size': 12, 'bold': True, 'alignment': 'center',
                    'space_before': 6, 'space_after': 6,
                },
                'content': {
                    'font_name_cn': '宋体', 'font_name_en': 'Times New Roman',
                    'font_size': 12, 'bold': False, 'alignment': 'center',
                    'vertical_alignment': 'center',
                },
                'style': {
                    'border_style': 'three_line',
                    'top_border_width': 1.5,
                    'middle_border_width': 0.75,
                    'bottom_border_width': 1.5,
                },
            },
            'recognition': {
                'table_caption_pattern': r'^表[0-9]+\.?[0-9]*\s+',
                'heading_patterns': {
                    'level1': r'^[0-9]+\s+',
                    'level2': r'^[0-9]+\.[0-9]+\s+',
                    'level3': r'^[0-9]+\.[0-9]+\.[0-9]+\s+',
                },
            },
        }

    def test_format_table_caption(self):
        self.doc.add_paragraph("表1.1 测试表格")
        fmt = TableFormatter(self.doc, self.config)
        fmt._format_table_captions()

        para = self.doc.paragraphs[0]
        self.assertTrue(para.runs[0].font.bold)

    def test_three_line_style(self):
        table = self.doc.add_table(rows=3, cols=3)
        table.cell(0, 0).text = "A"
        table.cell(1, 0).text = "B"
        table.cell(2, 0).text = "C"

        fmt = TableFormatter(self.doc, self.config)
        fmt._apply_three_line_style(table)


class TestCaptionFormatter(unittest.TestCase):
    def setUp(self):
        self.doc = Document()
        self.config = {
            'global': {'english_number_font': 'Times New Roman'},
            'figure': {
                'caption': {
                    'font_name_cn': '宋体', 'font_name_en': 'Times New Roman',
                    'font_size': 12, 'bold': True, 'alignment': 'center',
                    'space_before': 6, 'space_after': 6,
                },
            },
            'recognition': {
                'figure_caption_pattern': r'^图[0-9]+\.?[0-9]*\s+',
                'heading_patterns': {},
            },
        }

    def test_format_caption(self):
        self.doc.add_paragraph("图1 测试图片")
        fmt = CaptionFormatter(self.doc, self.config)
        fmt.format()

        para = self.doc.paragraphs[0]
        self.assertTrue(para.runs[0].font.bold)


class TestReferenceFormatter(unittest.TestCase):
    def setUp(self):
        self.doc = Document()
        self.config = {
            'global': {'english_number_font': 'Times New Roman'},
            'reference': {
                'title': {
                    'font_name_cn': '黑体', 'font_name_en': 'Times New Roman',
                    'font_size': 10.5, 'bold': True, 'alignment': 'center',
                    'space_before': 12, 'space_after': 6,
                },
                'content': {
                    'font_name_cn': '宋体', 'font_name_en': 'Times New Roman',
                    'font_size': 10.5, 'bold': False, 'alignment': 'left',
                    'line_spacing': 1.5,
                },
            },
            'recognition': {
                'reference_title_pattern': r'^参考文献$',
                'heading_patterns': {},
            },
        }

    def test_format_reference(self):
        self.doc.add_paragraph("参考文献")
        self.doc.add_paragraph("[1] 张三. 测试文献[J]. 测试期刊, 2024.")
        self.doc.add_paragraph("[2] Li S. Test Paper[J]. Test Journal, 2024.")

        fmt = ReferenceFormatter(self.doc, self.config)
        fmt.format()

        title_para = self.doc.paragraphs[0]
        self.assertTrue(title_para.runs[0].font.bold)
        self.assertEqual(title_para.runs[0].font.size.pt, 10.5)


class TestConfigLoading(unittest.TestCase):
    def test_load_default_config(self):
        config_path = os.path.join(os.path.dirname(__file__), '..', 'config', 'format_config.yaml')
        if os.path.exists(config_path):
            import yaml
            with open(config_path, 'r', encoding='utf-8') as f:
                config = yaml.safe_load(f)

            self.assertIn('global', config)
            self.assertIn('headings', config)
            self.assertIn('paragraph', config)
            self.assertIn('table', config)
            self.assertIn('figure', config)
            self.assertIn('reference', config)
            self.assertIn('toc', config)
            self.assertIn('recognition', config)
            self.assertEqual(config['global']['english_number_font'], 'Times New Roman')


if __name__ == '__main__':
    unittest.main()
